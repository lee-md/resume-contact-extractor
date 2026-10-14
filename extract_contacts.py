#!/usr/bin/env python3
"""
Contact Details Extractor
=========================
Reads resumes / profiles / contact lists (PDF, Word, text) and pulls out

    Name | Email | Designation | Contact Number

even when the document never writes "Name:" or "Phone:". The results are
written to an Excel (.xlsx) or CSV file. Everything runs offline - no AI,
no internet (except a one-time install of the helper packages).

How to run
----------
  python extract_contacts.py                         -> a window asks you to pick a folder
  python extract_contacts.py "C:\\Resumes"            -> every file in that folder (and sub-folders)
  python extract_contacts.py a.pdf b.docx -o out.xlsx
  python extract_contacts.py "C:\\Resumes" -o "C:\\My Sheet.xlsx"
        -> if My Sheet.xlsx already exists, rows are appended under ITS headers
           (e.g. "Full Name", "Mobile No", "Email ID", "Position" are recognised)

  python extract_contacts.py --help                  -> all options

Supported files: .pdf .docx .txt  (+ .doc .rtf .odt if MS Word or LibreOffice is installed)
                 scanned PDFs / images only if Tesseract OCR is installed (optional)
"""

import argparse
import csv
import importlib
import importlib.util
import os
import re
import shutil
import site
import subprocess
import sys
import tempfile
import traceback
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

# --------------------------------------------------------------------------
# Settings you may want to change
# --------------------------------------------------------------------------

# Column headers used when a NEW output file is created (order matters).
OUTPUT_COLUMNS = [
    ("name", "Name"),
    ("email", "Email"),
    ("designation", "Designation"),
    ("phone", "Contact Number"),
    ("file", "Source File"),
    ("notes", "Needs Review"),
]
DEFAULT_OUTPUT_NAME = "extracted_details.xlsx"
MAX_PHONES = 2  # how many different numbers to keep per person

# --------------------------------------------------------------------------
# One-time dependency install (so the user only needs plain Python)
# --------------------------------------------------------------------------

PACKAGES = [  # (import name(s), pip name)
    (("pymupdf", "fitz"), "pymupdf"),  # reads PDFs
    (("openpyxl",), "openpyxl"),        # writes .xlsx
]


def _have(*modules):
    return any(importlib.util.find_spec(m) is not None for m in modules)


def _pip_install(names):
    base = [sys.executable, "-m", "pip", "install", "--disable-pip-version-check", "-q"]
    for extra in ([], ["--user"]):
        try:
            subprocess.check_call(base + extra + list(names))
            return True
        except Exception:
            continue
    return False


def ensure_packages():
    missing = [pip for mods, pip in PACKAGES if not _have(*mods)]
    if not missing:
        return
    print("First run: installing helper packages (%s) ..." % ", ".join(missing))
    _pip_install(missing)
    # pymupdf has no build for some unusual Python versions -> pure-python fallback
    if not _have("pymupdf", "fitz") and not _have("pypdf"):
        _pip_install(["pypdf"])
    try:  # make packages installed with --user visible to this process
        site.addsitedir(site.getusersitepackages())
    except Exception:
        pass
    importlib.invalidate_caches()
    still = [pip for mods, pip in PACKAGES if not _have(*mods)]
    if still:
        print("WARNING: could not install %s. Try manually:  python -m pip install %s"
              % (", ".join(still), " ".join(still)))
    print()


def import_pymupdf():
    for name in ("pymupdf", "fitz"):
        try:
            return importlib.import_module(name)
        except Exception:
            pass
    return None


# --------------------------------------------------------------------------
# Reading files -> plain text (+ tables)
# --------------------------------------------------------------------------

class Document:
    def __init__(self, path, text="", tables=None, note=""):
        self.path = path
        self.text = text
        self._tables = tables  # list of tables, or a function returning them
        self.note = note

    @property
    def tables(self):
        if callable(self._tables):
            try:
                self._tables = self._tables()
            except Exception:
                self._tables = []
        return self._tables or []


TEXT_EXT = {".txt", ".text", ".md", ".csv"}
CONVERT_EXT = {".doc", ".rtf", ".odt", ".docm", ".dot"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
SUPPORTED = {".pdf", ".docx"} | TEXT_EXT | CONVERT_EXT | IMAGE_EXT


def read_document(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext in TEXT_EXT:
        with open(path, "rb") as f:
            raw = f.read()
        for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
            try:
                return Document(path, raw.decode(enc))
            except UnicodeDecodeError:
                continue
    if ext in CONVERT_EXT:
        return read_via_conversion(path)
    if ext in IMAGE_EXT:
        text = ocr_image(path)
        return Document(path, text, note="" if text else "image - install Tesseract OCR to read it")
    return Document(path, note="unsupported file type")


def read_pdf(path):
    fitz = import_pymupdf()
    if fitz is not None:
        with fitz.open(path) as pdf:
            text = "\n".join(page.get_text() for page in pdf)
        note = ""
        if len(text.strip()) < 30:
            text = ocr_pdf(path)
            note = "scanned PDF (OCR used)" if text.strip() else "no text found (scanned PDF? install Tesseract OCR)"

        def tables():
            out = []
            with fitz.open(path) as pdf:
                for page in pdf:
                    for t in page.find_tables().tables:
                        out.append([[c or "" for c in row] for row in t.extract()])
            return out

        return Document(path, text, tables, note)
    try:
        from pypdf import PdfReader
    except ImportError:
        return Document(path, note="cannot read PDF - run: python -m pip install pymupdf")
    reader = PdfReader(path)
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    return Document(path, text, note="" if text.strip() else "no text found (scanned PDF?)")


# --- .docx is a zip of XML files; read it with the standard library only ---
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _para_text(p):
    parts = []

    def walk(el):
        for ch in el:
            if ch.tag in (W + "p", MC_FALLBACK):
                continue  # nested paragraphs (text boxes) are read separately
            if ch.tag == W + "t":
                parts.append(ch.text or "")
            elif ch.tag == W + "tab":
                parts.append("  ")
            elif ch.tag in (W + "br", W + "cr"):
                parts.append("\n")
            else:
                walk(ch)

    walk(p)
    return "".join(parts)


def _docx_part(xml_bytes, lines, tables):
    root = ET.fromstring(xml_bytes)

    def walk(el):
        for ch in el:
            if ch.tag == MC_FALLBACK:  # duplicate copy of text boxes
                continue
            if ch.tag == W + "p":
                lines.append(_para_text(ch))
                walk(ch)  # text boxes anchored inside this paragraph
            elif ch.tag == W + "tbl":
                rows = []
                for tr in ch.findall(W + "tr"):
                    rows.append(["\n".join(_para_text(p) for p in tc.iter(W + "p"))
                                 for tc in tr.findall(W + "tc")])
                tables.append(rows)
                walk(ch)
            else:
                walk(ch)

    walk(root)


def read_docx(path):
    lines, tables = [], []
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            order = ([n for n in names if re.match(r"word/header\d*\.xml$", n)]
                     + ["word/document.xml"]
                     + [n for n in names if re.match(r"word/footer\d*\.xml$", n)])
            for name in order:
                if name in names:
                    _docx_part(z.read(name), lines, tables)
    except (zipfile.BadZipFile, KeyError, ET.ParseError):
        return Document(path, note="could not open Word file (corrupt or password protected?)")
    return Document(path, "\n".join(lines), tables)


def read_via_conversion(path):
    """Old .doc/.rtf/.odt: convert to .docx with LibreOffice or MS Word, then read."""
    tmp = tempfile.mkdtemp(prefix="extract_")
    try:
        out = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".docx")
        soffice = shutil.which("soffice") or next(
            (p for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                         r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                         "/Applications/LibreOffice.app/Contents/MacOS/soffice")
             if os.path.exists(p)), None)
        if soffice:
            subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, path],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        if not os.path.exists(out) and os.name == "nt":
            ps = ("$ErrorActionPreference='Stop'; $w=New-Object -ComObject Word.Application; "
                  "$w.Visible=$false; $w.DisplayAlerts=0; "
                  "try { $d=$w.Documents.Open($env:EX_SRC,$false,$true); $d.SaveAs2($env:EX_DST,16); "
                  "$d.Close($false) } finally { $w.Quit() }")
            env = dict(os.environ, EX_SRC=os.path.abspath(path), EX_DST=out)
            subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
                           env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        if os.path.exists(out):
            doc = read_docx(out)
            doc.path = path
            return doc
        return Document(path, note="old Word format - open it in Word and 'Save As' .docx (or install LibreOffice)")
    except Exception:
        return Document(path, note="could not convert file - save it as .docx or .pdf")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# --- optional OCR for scanned files (needs: pip install pytesseract pillow + Tesseract program) ---
def _tesseract():
    if not _have("PIL"):
        return None
    try:
        import pytesseract
    except ImportError:
        return None
    exe = shutil.which("tesseract") or next(
        (p for p in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                     r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe") if os.path.exists(p)), None)
    if not exe:
        return None
    pytesseract.pytesseract.tesseract_cmd = exe
    return pytesseract


def ocr_pdf(path):
    tess, fitz = _tesseract(), import_pymupdf()
    if not tess or not fitz:
        return ""
    from PIL import Image
    text = []
    with fitz.open(path) as pdf:
        for page in pdf:
            pix = page.get_pixmap(dpi=300)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            text.append(tess.image_to_string(img))
    return "\n".join(text)


def ocr_image(path):
    tess = _tesseract()
    if not tess:
        return ""
    from PIL import Image
    return tess.image_to_string(Image.open(path))


# --------------------------------------------------------------------------
# Vocabulary used by the heuristics
# --------------------------------------------------------------------------

JOB_WORDS = """
engineer modeler modeller developer designer architect analyst consultant manager director executive
officer lead head specialist coordinator co-ordinator administrator admin assistant associate intern
trainee technician supervisor foreman accountant auditor teacher professor lecturer tutor scientist
researcher programmer tester president vp founder co-founder ceo cto cfo coo cio cmo owner partner
representative agent clerk operator draftsman draughtsman drafter detailer estimator surveyor planner
inspector nurse doctor physician pharmacist chemist editor writer journalist photographer chef driver
receptionist secretary advisor adviser counsellor counselor recruiter buyer merchandiser strategist
economist lawyer advocate attorney banker cashier mechanic electrician plumber welder fitter artist
animator trainer instructor mentor coach principal dean controller superintendent steward attendant
expert freelancer contractor storekeeper storekeeper telecaller
""".split()
JOB_RE = re.compile(r"\b(?:%s)s?\b" % "|".join(re.escape(w) for w in JOB_WORDS), re.I)

# words that show a line is NOT a person's name
STOP_WORDS = set("""
resume cv curriculum vitae biodata bio-data profile summary summery objective career contact contacts
details personal information info education educational qualification qualifications academic
experience experiences work employment history professional skills skill technical software tools
languages language hobbies interests strengths declaration references reference projects project
certifications certification certificates courses achievements awards about me address phone mobile
email e-mail mail tel linkedin github website portfolio india indian male female married unmarried
single nationality gender date birth dob english hindi marathi tamil telugu kannada malayalam bengali
urdu gujarati punjabi present current till now page recap core competencies competency expertise key
responsibilities roles role and of the in at for with to from ms office excel word powerpoint autocad
revit navisworks autodesk bim mep hvac plumbing ltd pvt private limited inc llc llp company solutions
services technologies technology construction consultants university college school institute
academy polytechnic board b.tech btech b.e be m.tech diploma bachelor master degree ssc hsc sslc cgpa
gpa percentage street road nagar district dist city state pin pincode permanent communication
teamwork team leadership management problem solving aim good fluent basic maharashtra karnataka
kerala nadu telangana andhra pradesh gujarat rajasthan delhi mumbai pune bangalore bengaluru hyderabad
chennai kolkata ahmedabad noida gurgaon gurugram dubai uae usa uk qatar saudi oman thank thanks
signature place yours sincerely faithfully hereby confirm true knowledge belief passport visa
""".split())

# extra words for re-splitting text printed with spaced letters ("B I M  M O D E L E R")
SPLIT_VOCAB = set(JOB_WORDS) | {w for w in STOP_WORDS if len(w) >= 4} | set("""
bim mep civil senior junior software sales marketing project site design data structural mechanical
electrical quality safety hse account accounts finance human resources business development operations
""".split())

NOISE_FILENAME_WORDS = set("""
naukri resume cv curriculum vitae profile updated update latest final new copy doc docx pdf file my of
and the biodata bio data candidate application cover letter jan feb mar apr may jun jul aug sep oct nov dec
""".split())

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}")
LINKEDIN_RE = re.compile(r"linkedin\.com/in/([A-Za-z0-9\-_%]+)", re.I)
PHONE_RE = re.compile(r"(?<![\w+])(\+?\(?\d[\d \t().\-]{6,20}\d)(?!\w)")
PHONE_LABEL_RE = re.compile(r"phone|mobile|\bmob\b|\bph\b|\bcell\b|\btel\b|telephone|contact|whats\s?app|"
                            r"\bcall\b|\bm\s*:|\U0001F4DE|\u260E|\u2706|\U0001F4F1", re.I)
ID_LABEL_RE = re.compile(r"aadha+r|\bpan\b|passport|account|a/c|ifsc|\buan\b|\bpf\b|licen[cs]e|"
                         r"\bfax\b|\bgst|\bcin\b|\bid\b|roll|reg(istration)?\.?\s*no", re.I)
NAME_LABEL_RE = re.compile(r"^\s*(?:full\s*name|(?:candidate|applicant|employee|student)'?s?(?:\s*name)?|name)"
                           r"\s*[:\-\u2013]+\s*(.+)$", re.I)
DESIG_LABEL_RE = re.compile(
    r"^\s*(?:current\s+|present\s+)?(?:designation|job\s*title|position(?:\s+applied\s+for)?|"
    r"(?:job\s+|post\s+)?applied\s+for|job\s+applied\s+for|title|role|post|occupation)\s*[:\-\u2013]*\s*(.*)$", re.I)
SENTENCE_RE = re.compile(r"\b(?:years?|experience|experienced|seeking|looking|responsible|worked|working|"
                         r"i|am|my|we|our|will|have|has|using|including|ensure|ensuring|with)\b", re.I)
COMPANY_RE = re.compile(r"^the\b|\b(?:pvt|private|ltd|limited|inc|llc|llp|corp|corporation|company|"
                        r"consultancy|consulting|industries|enterprises|associates|engineers|enginners|"
                        r"consultants|developers|designers|architects|builders|contractors|planners|"
                        r"surveyors)\b|\band\s+consultant\b", re.I)
PARA_HEAD_RE = re.compile(r"^\W*(?:aim|objective|career\s+objective|summary|professional\s+summary|profile|"
                          r"about(?:\s+me)?)\W*$", re.I)
CURRENT_RE = re.compile(r"\b(?:present|current|currently|till\s+(?:date|now)|to\s+date|now)\b", re.I)
EDUCATION_RE = re.compile(r"bachelor|master|diploma|b\.?\s?tech|m\.?\s?tech|\bb\.?e\b|degree|university|"
                          r"college|school|institute|course|certificat", re.I)


# --------------------------------------------------------------------------
# Text clean-up helpers
# --------------------------------------------------------------------------

def normalize_text(text):
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[\u200b-\u200f\u2060\ufeff\u00ad]", "", text)
    # private-use characters are icon fonts (phone/mail symbols) -> space
    text = "".join(" " if unicodedata.category(c) == "Co" else c for c in text)
    text = text.replace("\t", "  ").replace("\r", "\n")
    # e-mail broken over two lines: "name123\n@gmail.com"
    text = re.sub(r"([A-Za-z0-9._%+\-])[ ]*\n[ ]*(@[A-Za-z0-9\-]+\.)", r"\1\2", text)
    text = re.sub(r"(@[A-Za-z0-9\-]+)[ ]*\n[ ]*(\.[A-Za-z]{2,}\b)", r"\1\2", text)
    return text


def segment_glued(word, vocab):
    """'BIMMODELER' -> 'BIM MODELER' using known words; unchanged unless fully covered."""
    low, n = word.lower(), len(word)
    best = [None] * (n + 1)  # best[i] = list of cut positions covering word[:i]
    best[0] = []
    for i in range(n):
        if best[i] is None:
            continue
        for j in range(i + 3, n + 1):
            if low[i:j] in vocab and (best[j] is None or len(best[i]) + 1 < len(best[j])):
                best[j] = best[i] + [j]
    if not best[n] or len(best[n]) < 2:
        return word
    cuts, out, start = best[n], [], 0
    for c in cuts:
        out.append(word[start:c])
        start = c
    return " ".join(out)


def fix_spaced_letters(line, vocab):
    """'J O H N  S M I T H' -> 'JOHN SMITH'."""
    s = line.strip()
    if not re.fullmatch(r"(?:[^\W\d_&]\s{1,3}){3,}[^\W\d_]", s):
        return line
    words = ["".join(w.split()) for w in re.split(r"\s{2,}", s)]
    return " ".join(segment_glued(w, vocab) if len(words) == 1 else w for w in words)


def tidy(s):
    return re.sub(r"\s+", " ", s).strip()


def split_camel(s):
    return re.findall(r"[A-Z]?[a-z]+|[A-Z]+(?![a-z])", s)


# --------------------------------------------------------------------------
# Field finders
# --------------------------------------------------------------------------

def find_emails(text):
    seen, out = set(), []
    for m in EMAIL_RE.finditer(text):
        e = m.group(0).strip(".").lower()
        if e not in seen and not e.endswith((".png", ".jpg", ".jpeg", ".gif")):
            seen.add(e)
            out.append(e)
    return out


def _valid_phone(raw):
    digits = re.sub(r"\D", "", raw)
    if not 10 <= len(digits) <= 13 or len(set(digits)) <= 2:
        return None
    if re.search(r"(?:19|20)\d\d\s*[-\u2013/]\s*(?:19|20)?\d\d", raw):  # 2014-2018, dates
        return None
    groups = re.findall(r"\d+", raw)
    if len(groups) > 1 and all(re.fullmatch(r"(?:19|20)\d\d", g) for g in groups):
        return None
    num = tidy(raw)
    if num.count("(") != num.count(")"):
        num = num.replace("(", "").replace(")", "")
    return num.strip(" -.")


def find_phones(lines):
    """Returns [(line_index, number, labelled)] in document order."""
    found = []
    for i, line in enumerate(lines):
        if not line:
            continue
        clean = LINKEDIN_RE.sub(" ", EMAIL_RE.sub(" ", line))
        clean = re.sub(r"https?://\S+|www\.\S+", " ", clean)
        for m in PHONE_RE.finditer(clean):
            context = clean[max(0, m.start() - 25):m.start()]
            if ID_LABEL_RE.search(context):
                continue
            raw = m.group(1)
            nums = [_valid_phone(raw)]
            if not nums[0]:  # e.g. two numbers glued by spaces: try the pieces
                nums = [_valid_phone(p) for p in re.split(r"\s{2,}|\s[-\u2013]\s", raw)]
            prev = lines[i - 1] if i else ""
            labelled = bool(PHONE_LABEL_RE.search(context) or
                            (PHONE_LABEL_RE.search(prev) and len(prev) < 25) or
                            (i + 1 < len(lines) and PHONE_LABEL_RE.fullmatch(lines[i + 1].strip(" :"))))
            for num in nums:
                if num:
                    found.append((i, num, labelled))
    return found


class Hints:
    """Clues about the person's name from the e-mail, LinkedIn URL and file name."""

    def __init__(self, path, text, emails):
        self.email = " ".join(re.sub(r"[^a-z]", "", e.split("@")[0]) for e in emails)
        slug = LINKEDIN_RE.search(text.replace("\n", ""))
        self.linkedin_tokens = []
        if slug:
            self.linkedin_tokens = [t.lower() for t in re.split(r"[-_%]+", slug.group(1))
                                    if t.isalpha() and len(t) >= 2]
        stem = os.path.splitext(os.path.basename(path))[0] if path else ""
        stem = re.sub(r"\[.*?\]|\(.*?\)", " ", stem)
        words = []
        for chunk in re.split(r"[^A-Za-z]+", stem):
            words.extend(split_camel(chunk) if chunk else [])
        self.file_tokens = [w.lower() for w in words if w.lower() not in NOISE_FILENAME_WORDS and len(w) >= 2]
        self.file_words = [w for w in words if w.lower() not in NOISE_FILENAME_WORDS and len(w) >= 2]

    def sources(self, exclude=None):
        out = {"email": self.email, "linkedin": "".join(self.linkedin_tokens),
               "file": "".join(self.file_tokens)}
        out.pop(exclude, None)
        return [v for v in out.values() if v]

    def vocab(self):
        return {t for t in self.file_tokens + self.linkedin_tokens if len(t) >= 3}

    def overlap(self, name, exclude=None):
        """How many words of `name` also appear in the e-mail / LinkedIn / file name."""
        srcs = self.sources(exclude)
        return sum(1 for t in re.findall(r"[a-z]+", name.lower())
                   if len(t) >= 3 and any(t in s for s in srcs))


def clean_for_name(raw):
    s = raw.strip()
    m = NAME_LABEL_RE.match(s)
    if m:
        s = m.group(1)
    s = re.sub(r"^[\W_]+|[^\w.]+$", "", s)
    return tidy(s)


def looks_like_name(s):
    tokens = s.split()
    if not 1 <= len(tokens) <= 5 or len(s) > 45:
        return False
    if not all(re.fullmatch(r"\.?[^\W\d_][^\W\d_'\u2019\-]*\.?(?:[^\W\d_]\.?)*", t) for t in tokens):
        return False
    if sum(c.isalpha() for c in s) < 3 or all(len(t.strip(".")) == 1 for t in tokens):
        return False
    if not s[0].isupper() and not s.lstrip(".")[0:1].isupper():
        return False
    if any(t.lower().strip(".") in STOP_WORDS for t in tokens) or JOB_RE.search(s):
        return False
    return True


def is_designation(s):
    s = re.sub(r"^[\W\d_]*(?=[^\W\d_])", "", s.strip())  # bullets / "2." numbering
    if not s or len(tidy(s)) > 120 or "@" in s or re.search(r"https?:|www\.", s, re.I):
        return False
    if s[0].islower() or s.count(",") >= 2 or re.search(r"(?:,|&|\band|\bof)$", s.rstrip(), re.I):
        return False  # piece of a sentence or a comma-separated skills list
    if re.search(r"\d{5,}", s):
        return False
    core = clean_designation(s)
    if not JOB_RE.search(core) or len(core.split()) > 10:
        return False
    return not (EDUCATION_RE.search(core) or SENTENCE_RE.search(core))


def clean_designation(s, name=""):
    """'BIM Engineer \u2013 MEP | Revit & Navisworks' -> 'BIM Engineer \u2013 MEP' (drops dates, companies, places)."""
    s = tidy(re.sub(r"\s{3,}", " | ", s.strip()))
    if name:
        s = re.sub(re.escape(name), " ", s, flags=re.I)
    m = DESIG_LABEL_RE.match(s)
    if m and m.group(1):
        s = m.group(1)
    keep = []
    for part in re.split(r"\s*[|:\u2022\u25cf]\s*|\s+at\s+|,\s+", s):
        bits = re.split(r"(\s+[-\u2013\u2014]\s+|(?<=[a-z.])[-\u2013\u2014]\s+)", part.strip(" -\u2013:,[]"))
        # keep "BIM Engineer \u2013 MEP" but drop "\u2013 Acme Industries Pvt. Ltd." / "XYZ Engineers \u2013" / "\u2013 Jan 2024"
        text = ""
        for k in range(0, len(bits), 2):
            bit = bits[k]
            if bit and not COMPANY_RE.search(bit) and not (k and re.search(r"\d", bit)):
                text += (bits[k - 1] if text else "") + bit
        if text and JOB_RE.search(text) and not COMPANY_RE.search(text):
            keep.append(text)
    s = " | ".join(keep)
    s = re.sub(r"^[^\w(\[]+|[^\w)\]]+$", "", s)
    if (s.startswith("[") and s.endswith("]")) or (s.startswith("(") and s.endswith(")")):
        s = s[1:-1]
    return tidy(s)


def find_name(lines, hints, contact_idx):
    """Score every plausible line and return (name, line_index, score, source)."""
    nonempty = [i for i, l in enumerate(lines) if l.strip()]
    rank = {i: r for r, i in enumerate(nonempty)}
    cands = []

    def score(text, i, n_lines=1):
        tokens = text.split()
        sc = 1.0
        r = rank.get(i, 99)
        sc += 3 if r == 0 else 2.5 if r < 3 else 2 if r < 6 else 1 if r < 12 else 0
        sc += 3 * min(hints.overlap(text), 3)
        if any(abs(i - c) <= 4 for c in contact_idx):
            sc += 1.5
        following = [lines[j] for j in range(i + n_lines, min(len(lines), i + n_lines + 3)) if lines[j].strip()]
        if following and is_designation(following[0]):
            sc += 1.5
        sc += 1 if 2 <= len(tokens) <= 4 else -1.5 if len(tokens) == 1 else 0
        if text.isupper() or all(t[:1].isupper() for t in tokens if t[0].isalpha()):
            sc += 0.5
        if n_lines > 1:
            sc -= 0.5
        return sc

    for i in nonempty:
        raw = lines[i]
        labelled = bool(NAME_LABEL_RE.match(raw))
        whole = clean_for_name(raw)
        options = [whole] + [clean_for_name(p) for p in re.split(r"\s+[|\u2013\-,]\s+|\s{3,}", whole)[:1]]
        if ":" in raw:  # "Some Label: Firstname Lastname"
            options.append(clean_for_name(raw.split(":", 1)[1]))
        for opt in dict.fromkeys(options):
            if looks_like_name(opt):
                cands.append((score(opt, i) + (12 if labelled else 0), i, opt, "text"))
        # name broken over two lines ("RAVI KUMAR" / "SHARMA")
        nxt = next((j for j in range(i + 1, min(len(lines), i + 3)) if lines[j].strip()), None)
        if nxt is not None and looks_like_name(whole):
            two = clean_for_name(lines[nxt])
            if looks_like_name(two) and len((whole + " " + two).split()) <= 4:
                joined = whole + " " + two
                cands.append((score(joined, i, nxt - i + 1), i, joined, "text"))

    # fall-backs: LinkedIn URL and file name
    if len(hints.linkedin_tokens) >= 2:
        nm = " ".join(t.capitalize() for t in hints.linkedin_tokens[:4])
        cands.append((2 + 3 * min(hints.overlap(nm, "linkedin"), 3), -1, nm, "LinkedIn URL"))
    if 2 <= len(hints.file_words) <= 4:
        nm = " ".join(w.capitalize() for w in hints.file_words)
        cands.append((0.5 + 3 * min(hints.overlap(nm, "file"), 3), -1, nm, "file name"))

    if not cands:
        return "", -1, 0, ""
    best = max(cands, key=lambda c: (c[0], -c[1] if c[1] >= 0 else -10 ** 6))
    sc, i, nm, src = best
    if nm.isupper() or nm.islower():
        nm = nm.title()
    return nm, i, sc, src


def find_designation(lines, name_idx, name):
    cands = []
    for i, raw in enumerate(lines):
        line = raw.strip()
        if not line:
            continue
        m = DESIG_LABEL_RE.match(line)
        if m and not line.lower().startswith(("roles", "role &", "role and")):
            value = m.group(1).strip()
            if not value:  # label on its own line -> value on the next line
                value = next((lines[j].strip() for j in range(i + 1, min(len(lines), i + 3))
                              if lines[j].strip()), "")
            if value and JOB_RE.search(value) and len(value) <= 90:
                cands.append((20, i, value))
                continue
        if not is_designation(line):
            continue
        # wrapped paragraph text ("AIM:" / "Civil Engineer and BIM" / "Modeler with 2 years ...")
        prev = next((lines[j].strip() for j in range(i - 1, -1, -1) if lines[j].strip()), "")
        nxt = next((lines[j].strip() for j in range(i + 1, len(lines)) if lines[j].strip()), "")
        if (SENTENCE_RE.search(nxt) and not line.endswith((".", "!", "?"))
                and (PARA_HEAD_RE.match(prev) or prev.endswith(":") or SENTENCE_RE.search(prev))):
            continue
        if name_idx >= 0 and 0 < i - name_idx <= 3:
            sc = 10 - (i - name_idx)
        elif name_idx >= 0 and i == name_idx and name and len(clean_designation(line, name)) > 2:
            sc = 9  # "John Smith | Software Engineer"
        else:
            sc = 4 if i < 15 else 1
            sc += 1 if len(clean_designation(line).split()) <= 5 else 0
            sc += 2 if CURRENT_RE.search(line) else 0  # "BIM Modeler | Dec 2024 - Present"
        cands.append((sc, i, line))
    if not cands:
        return ""
    best = max(cands, key=lambda c: (c[0], -c[1]))
    return clean_designation(best[2], name)


def extract_record(lines, hints):
    """Extract one person's details from a list of text lines."""
    text = "\n".join(lines)
    emails = find_emails(text)
    phones = find_phones(lines)
    contact_idx = sorted({i for i, l in enumerate(lines) if EMAIL_RE.search(l)} | {p[0] for p in phones})
    name, name_idx, name_score, name_src = find_name(lines, hints, contact_idx)
    designation = find_designation(lines, name_idx, name)

    # primary e-mail: the one that matches the name best, else the first
    email = ""
    if emails:
        name_tokens = [t for t in re.findall(r"[a-z]+", name.lower()) if len(t) >= 3]
        email = max(emails, key=lambda e: (sum(t in e.split("@")[0] for t in name_tokens), -emails.index(e)))

    # phones: labelled ones first, de-duplicated on the last 10 digits
    numbers, seen = [], set()
    for _, num, _ in sorted(phones, key=lambda p: (not p[2], p[0])):
        key = re.sub(r"\D", "", num)[-10:]
        if key not in seen:
            seen.add(key)
            numbers.append(num)

    notes = []
    if not name:
        notes.append("name not found")
    elif name_src != "text":
        notes.append("name taken from %s" % name_src)
    elif name_score < 6 or not (emails or numbers):
        notes.append("check name")
    if not email:
        notes.append("no email")
    if not numbers:
        notes.append("no phone")
    if not designation:
        notes.append("no designation")
    return {"name": name, "email": email, "designation": designation,
            "phone": ", ".join(numbers[:MAX_PHONES]), "notes": "; ".join(notes)}


# --------------------------------------------------------------------------
# Files that contain MANY people (contact lists, tables)
# --------------------------------------------------------------------------

HEADER_SYNONYMS = {
    "name": ["name", "full name", "candidate name", "candidate", "person", "contact person", "contact name",
             "employee name", "employee", "applicant", "applicant name", "student name", "member name"],
    "email": ["email", "e mail", "email id", "emailid", "email address", "mail", "mail id"],
    "designation": ["designation", "title", "job title", "position", "role", "post", "current designation",
                    "job role", "occupation", "profile"],
    "phone": ["contact number", "contact no", "contact", "phone", "phone number", "phone no", "mobile",
              "mobile number", "mobile no", "cell", "telephone", "tel", "whatsapp", "number", "contact details"],
    "file": ["source file", "file", "file name", "filename", "source", "document", "resume file"],
    "notes": ["needs review", "notes", "remarks", "comments", "review", "status"],
    "serial": ["sr no", "sr", "s no", "sl no", "serial", "serial no", "serial number", "no", "sno", "srno"],
}


def header_field(header):
    if str(header or "").strip() == "#":
        return "serial"
    h = re.sub(r"[^a-z]+", " ", str(header or "").lower()).strip()
    if not h:
        return None
    for field, syns in HEADER_SYNONYMS.items():
        if h in syns:
            return field
    words = set(h.split())
    if words & {"email", "mail"}:
        return "email"
    if words & {"name", "person"}:
        return "name"
    if words & {"phone", "mobile", "contact", "cell", "telephone", "tel", "whatsapp"}:
        return "phone"
    if words & {"designation", "title", "position", "role", "occupation"}:
        return "designation"
    return None


def records_from_table(rows, hints):
    rows = [[tidy(str(c or "")) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        return []
    header = {i: header_field(c) for i, c in enumerate(rows[0])}
    header = {i: f for i, f in header.items() if f in ("name", "email", "designation", "phone")}
    data = rows[1:] if len(header) >= 2 else rows
    if len(header) < 2:
        header = {}
    out = []
    for row in data:
        rec = {}
        for i, cell in enumerate(row):
            f = header.get(i)
            if f == "email":
                em = find_emails(cell)
                rec["email"] = em[0] if em else ""
            elif f == "phone":
                rec["phone"] = ", ".join(n for _, n, _ in find_phones([cell])[:MAX_PHONES]) or cell
            elif f:
                rec[f] = cell
        for cell in row:  # no (useful) header -> guess each cell from its content
            if not rec.get("email") and find_emails(cell):
                rec["email"] = find_emails(cell)[0]
            elif not rec.get("phone") and find_phones([cell]):
                rec["phone"] = ", ".join(n for _, n, _ in find_phones([cell])[:MAX_PHONES])
            elif not rec.get("designation") and is_designation(cell):
                rec["designation"] = clean_designation(cell)
            elif not rec.get("name") and looks_like_name(clean_for_name(cell)):
                rec["name"] = clean_for_name(cell)
        if rec.get("email") or rec.get("phone"):
            missing = [f for f in ("name", "email", "phone", "designation") if not rec.get(f)]
            rec["notes"] = "; ".join("no " + f for f in missing)
            out.append(rec)
    return out


RESUME_HEADINGS = re.compile(r"^\W*(?:work\s+|professional\s+)?(?:experience|education|skills|"
                             r"summary|objective|projects|declaration)\b", re.I | re.M)


def split_blocks(lines):
    """Split a contact list into one block of lines per person (one e-mail each)."""
    first_seen, email_lines = set(), []
    for i, l in enumerate(lines):
        for e in find_emails(l):
            if e not in first_seen:
                first_seen.add(e)
                email_lines.append(i)
    email_lines = sorted(set(email_lines))
    if len(email_lines) < 2:
        return [lines]
    starts = [0]
    for a, b in zip(email_lines, email_lines[1:]):
        start = b
        for j in range(a + 1, b + 1):
            if looks_like_name(clean_for_name(lines[j])):
                start = j
                break
        starts.append(start)
    starts.append(len(lines))
    return [lines[s:e] for s, e in zip(starts, starts[1:]) if s < e]


def extract_from_document(doc, mode):
    text = normalize_text(doc.text or "")
    if not text.strip():
        return [{"name": "", "email": "", "designation": "", "phone": "",
                 "notes": doc.note or "no text found"}]
    hints = Hints(doc.path, text, find_emails(text))
    vocab = SPLIT_VOCAB | hints.vocab()
    lines = [fix_spaced_letters(l, vocab).strip() for l in text.split("\n")]

    n_emails = len(find_emails(text))
    looks_like_resume = len(set(m.lower() for m in RESUME_HEADINGS.findall(text))) >= 2
    if mode == "multi" or (mode == "auto" and n_emails >= 2 and not (looks_like_resume and n_emails < 5)):
        recs = []
        for table in doc.tables:
            recs.extend(records_from_table(table, hints))
        if len(recs) < 2:
            recs = [extract_record(block, hints) for block in split_blocks(lines)]
            recs = [r for r in recs if r["email"] or r["phone"]]
        if len(recs) >= 2 or mode == "multi":
            for r in recs:
                r["notes"] = "; ".join(x for x in (r.get("notes"), doc.note) if x)
            return recs
    rec = extract_record(lines, hints)
    if doc.note:
        rec["notes"] = "; ".join(x for x in (rec["notes"], doc.note) if x)
    return [rec]


# --------------------------------------------------------------------------
# Output (Excel / CSV)
# --------------------------------------------------------------------------

def _existing_headers(path):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return next(csv.reader(f), [])
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).active
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return [c if c is not None else "" for c in row]


def _existing_column(path, col):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [r[col] for r in list(csv.reader(f))[1:] if len(r) > col]
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).active
    return [r[col] for r in ws.iter_rows(min_row=2, values_only=True) if len(r) > col]


def _last_serial(path, col):
    nums = [int(float(v)) for v in _existing_column(path, col)
            if re.fullmatch(r"\s*\d+(\.0)?\s*", str(v if v is not None else ""))]
    return max(nums) if nums else 0


def write_output(records, path, allow_duplicates=False):
    """Writes/append rows. If the file exists its own header row decides the columns."""
    if path.lower().endswith(".xlsx") and not _have("openpyxl"):
        path = os.path.splitext(path)[0] + ".csv"
        print("openpyxl is not installed -> saving as CSV instead: %s" % path)

    exists = os.path.exists(path) and os.path.getsize(path) > 0
    headers = _existing_headers(path) if exists else []
    if not any(str(h).strip() for h in headers):
        headers, exists_with_rows = [h for _, h in OUTPUT_COLUMNS], False
    else:
        exists_with_rows = True
    fields = [header_field(h) for h in headers]
    if exists_with_rows:
        unmatched = [h for h, f in zip(headers, fields) if h and not f]
        print("Appending to existing file. Matched columns: %s" %
              ", ".join("%s->%s" % (h, f) for h, f in zip(headers, fields) if f))
        if unmatched:
            print("  (left blank: %s)" % ", ".join(map(str, unmatched)))
        if "email" not in fields and "name" not in fields:
            print("  WARNING: could not recognise any Name/Email column in that file's first row.")

    skipped = []
    if not allow_duplicates:  # same e-mail already in the sheet, or twice in this run
        known = set()
        if exists_with_rows and "email" in fields:
            known = {str(v or "").strip().lower() for v in _existing_column(path, fields.index("email"))}
        unique = []
        for r in records:
            e = (r.get("email") or "").lower()
            if e and e in known:
                skipped.append(r)
                continue
            if e:
                known.add(e)
            unique.append(r)
        records = unique

    serial = _last_serial(path, fields.index("serial")) if exists_with_rows and "serial" in fields else 0
    rows = []
    for n, r in enumerate(records, serial + 1):
        rows.append([n if f == "serial" else r.get(f, "") if f else "" for f in fields])

    if path.lower().endswith(".csv"):
        with open(path, "a" if exists_with_rows else "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            if not exists_with_rows:
                w.writerow(headers)
            w.writerows(rows)
    else:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        if exists_with_rows:
            wb = openpyxl.load_workbook(path)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Extracted"
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="305496")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
        yellow = PatternFill("solid", fgColor="FFF2CC")
        notes_col = fields.index("notes") if "notes" in fields else None
        phone_col = fields.index("phone") + 1 if "phone" in fields else None
        for row in rows:
            ws.append(row)
            r = ws.max_row
            if phone_col:
                ws.cell(r, phone_col).number_format = "@"  # keep phone numbers as text
            if notes_col is not None and row[notes_col]:
                for c in ws[r]:
                    c.fill = yellow
        if not exists_with_rows:
            for i, h in enumerate(headers, 1):
                width = max([len(str(h))] + [len(str(rw[i - 1])) for rw in rows]) + 2
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(width, 60)
        wb.save(path)
    return path, len(rows), skipped


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def collect_files(inputs, recursive=True):
    files = []
    for p in inputs:
        p = os.path.abspath(p.strip('"'))
        if os.path.isdir(p):
            for root, dirs, names in os.walk(p):
                for n in sorted(names):
                    if os.path.splitext(n)[1].lower() in SUPPORTED and not n.startswith("~$"):
                        files.append(os.path.join(root, n))
                if not recursive:
                    break
        elif os.path.isfile(p):
            files.append(p)
        else:
            print("Not found: %s" % p)
    return files


def pick_folder():
    try:
        import tkinter
        from tkinter import filedialog
        root = tkinter.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        folder = filedialog.askdirectory(title="Choose the folder that contains the files (PDF / Word)")
        root.destroy()
        return folder
    except Exception:
        return input("Folder that contains the files: ").strip().strip('"')


def main(argv=None):
    ap = argparse.ArgumentParser(description="Extract Name / Email / Designation / Contact Number from "
                                             "PDF, Word and text files into Excel or CSV.")
    ap.add_argument("inputs", nargs="*", help="files and/or folders (default: a folder picker opens)")
    ap.add_argument("-o", "--output", help="output .xlsx or .csv (default: %s inside the input folder). "
                                           "If it exists, rows are appended under its headers." % DEFAULT_OUTPUT_NAME)
    ap.add_argument("--mode", choices=["auto", "single", "multi"], default="auto",
                    help="single = one person per file (resumes); multi = many people per file "
                         "(contact lists); auto = decide per file (default)")
    ap.add_argument("--no-subfolders", action="store_true", help="do not look inside sub-folders")
    ap.add_argument("--allow-duplicates", action="store_true",
                    help="when appending, also add people whose e-mail is already in the sheet")
    ap.add_argument("--show-text", action="store_true", help="print the text read from each file (debugging)")
    args = ap.parse_args(argv)

    ensure_packages()

    inputs = args.inputs
    if not inputs:
        folder = pick_folder()
        if not folder:
            print("No folder chosen.")
            return 1
        inputs = [folder]

    files = collect_files(inputs, recursive=not args.no_subfolders)
    out_path = args.output
    if out_path:
        out_path = os.path.abspath(out_path)
        if os.path.splitext(out_path)[1].lower() not in (".xlsx", ".csv"):
            out_path += ".xlsx"
    else:
        first = os.path.abspath(inputs[0].strip('"'))
        base = first if os.path.isdir(first) else os.path.dirname(first)
        out_path = os.path.join(base, DEFAULT_OUTPUT_NAME)
    files = [f for f in files if os.path.abspath(f) != out_path]
    if not files:
        print("No supported files found (%s)." % ", ".join(sorted(SUPPORTED)))
        return 1

    print("Reading %d file(s)...\n" % len(files))
    records, errors = [], {}
    for n, path in enumerate(files, 1):
        name = os.path.basename(path)
        try:
            doc = read_document(path)
            if args.show_text:
                print("----- %s -----\n%s\n" % (name, normalize_text(doc.text)))
            recs = extract_from_document(doc, args.mode)
        except Exception as e:
            errors[path] = traceback.format_exc()
            recs = [{"name": "", "email": "", "designation": "", "phone": "",
                     "notes": "error reading file: %s" % e}]
        for r in recs:
            r["file"] = name
            r["path"] = path
        records.extend(recs)
        first = recs[0]
        more = "  (+%d more people)" % (len(recs) - 1) if len(recs) > 1 else ""
        print("[%*d/%d] %s" % (len(str(len(files))), n, len(files), name))
        print("        %s | %s | %s | %s%s" % (first["name"] or "-", first["email"] or "-",
                                               first["phone"] or "-", first["designation"] or "-", more))
        if first.get("notes"):
            print("        check: %s" % first["notes"])

    while True:
        try:
            saved, written, skipped = write_output(records, out_path, args.allow_duplicates)
            break
        except PermissionError:
            ans = input("\nCannot write %s - is it open in Excel? Close it and press Enter "
                        "(or type q to quit): " % out_path)
            if ans.strip().lower() == "q":
                return 1

    review = sum(1 for r in records if r.get("notes") and r["path"] not in errors)
    print("\nDone. %d row(s) written to:\n  %s" % (written, saved))
    if skipped:
        print("  %d skipped as duplicates (same e-mail already in the sheet or seen twice)." % len(skipped))
    if review:
        print("  %d row(s) have something to double-check (see the 'Needs Review' column, "
              "highlighted yellow in Excel)." % review)
    if errors:
        print("  %d file(s) could not be read at all." % len(errors))
    log_path = os.path.splitext(saved)[0] + "_log.txt"
    try:
        write_log(log_path, saved, files, records, skipped, errors)
        print("  Log of what needs checking: %s" % log_path)
    except OSError as e:
        print("  Could not write the log file: %s" % e)
    return 0


def write_log(log_path, output, files, records, skipped, errors):
    """Plain-text report: which files need a manual check and why, errors, duplicates."""
    import datetime

    def summary(r):
        return " | ".join(r.get(f) or "-" for f in ("name", "email", "phone", "designation"))

    review = [r for r in records if r.get("notes") and r.get("path") not in errors]
    skipped_ids = {id(r) for r in skipped}
    out = ["Contact extraction log - %s" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
           "Output file : %s" % output,
           "Files read  : %d   rows found: %d   need review: %d   errors: %d   duplicates skipped: %d"
           % (len(files), len(records), len(review), len(errors), len(skipped)),
           "", "Rows marked REVIEW were still written to the sheet - only the listed part is missing or",
           "uncertain. Rows marked OK were found confidently, but it is still a guess: spot-check a few.", ""]

    if errors:
        out += ["=" * 70, "ERRORS - these files could not be read at all (damaged, password-protected",
                "or not really the file type its name says). Open them yourself to check.", "=" * 70]
        for path, tb in errors.items():
            out += [path, "    reason  : %s" % tb.strip().splitlines()[-1], ""]
    if review:
        out += ["=" * 70, "NEEDS REVIEW - open these files and check the sheet", "=" * 70]
        for r in review:
            out += [r["path"], "    problem : %s" % r["notes"], "    found   : %s" % summary(r), ""]
    if skipped:
        out += ["=" * 70, "DUPLICATES - not written (same e-mail already in the sheet or seen twice)",
                "=" * 70]
        for r in skipped:
            out += ["%s  ->  %s" % (r["file"], r.get("email"))]
        out.append("")

    out += ["=" * 70, "ALL FILES", "=" * 70]
    for r in records:
        status = ("ERROR" if r.get("path") in errors else "SKIPPED" if id(r) in skipped_ids
                  else "REVIEW" if r.get("notes") else "OK")
        out.append("%-7s %s  ->  %s" % (status, r["file"], summary(r)))

    with open(log_path, "w", encoding="utf-8-sig") as f:
        f.write("\n".join(out) + "\n")


if __name__ == "__main__":
    interactive = len(sys.argv) == 1 or os.environ.get("EXTRACTOR_PAUSE") == "1"
    try:
        code = main()
    except KeyboardInterrupt:
        code = 1
    except Exception:
        traceback.print_exc()
        code = 1
    if interactive:
        try:
            input("\nPress Enter to close...")
        except EOFError:
            pass
    sys.exit(code)
